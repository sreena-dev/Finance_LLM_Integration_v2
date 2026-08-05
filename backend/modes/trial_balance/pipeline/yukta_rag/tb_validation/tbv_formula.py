"""Layer 4 — formula-cell flag (presence-only, always attempted).

Requires raw formula-text access, which is NOT available anywhere else in
this codebase's ingestion pipeline today (confirmed: every existing reader —
``pandas.read_excel``, ``pandas.read_html`` — only ever exposes computed
values, never formula strings; no ``openpyxl.load_workbook(...,
data_only=False)`` call exists elsewhere). This is a genuine prerequisite
gap this module closes for itself, not something silently skipped.

Only meaningful for the raw-bytes ingestion path (``ingest_from_bytes``) —
the doc_id/convenience path has no raw file to re-read, so callers pass
``data=None`` there and get an explicit "unavailable" result instead of an
empty-but-silent list.
"""

from __future__ import annotations

import io

import openpyxl

from yukta_rag.tb_validation.tbv_ingest import ParsedTBTable

_ROLE_COLUMNS = ("opening_balance", "closing_balance", "debit", "credit")


def scan_formulas(data: bytes | None, table: ParsedTBTable) -> list[dict] | None:
    """Scan the given raw workbook bytes for formula cells in the
    opening/debit/credit/closing columns, using the SAME column positions
    Step-0 already resolved on this table. Returns None (not []) when raw
    bytes aren't available at all — the caller must treat None as
    "unavailable", not "checked, found nothing"."""
    if data is None or not table.has_raw_access or table.header_row is None:
        return None

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    sheet = (wb[table.sheet_name] if table.sheet_name and table.sheet_name in wb.sheetnames
            else wb.worksheets[0])

    flags: list[dict] = []
    code_col = table.role_cols.get("ledger_code")
    for row_idx, row in enumerate(sheet.iter_rows(min_row=table.header_row + 2)):  # +2: 1-based, skip header
        for role in _ROLE_COLUMNS:
            col = table.role_cols.get(role)
            if col is None or col >= len(row):
                continue
            cell = row[col]
            raw = cell.value
            if isinstance(raw, str) and raw.startswith("="):
                label_cell = row[code_col] if code_col is not None and code_col < len(row) else None
                flags.append({
                    "row": table.header_row + 2 + row_idx,
                    "account_label": (label_cell.value if label_cell is not None else None),
                    "column": role,
                    "value": cell.value,
                    "formula": raw,
                    "message": (f"{role} cell contains a live formula ({raw}) rather than a "
                               "static value — purely informational, not an assessment of "
                               "the value itself."),
                    "severity": "info",
                })
    wb.close()
    return flags
