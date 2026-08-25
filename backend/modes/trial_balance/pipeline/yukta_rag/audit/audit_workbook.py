"""Generate the downloadable TB_Audit.xlsx workbook for Audit mode.

Six sheets, all deterministic (no LLM):
  1. Financial Snapshot   — key metrics per period
  2. Snapshot Drilldown   — EVERY metric broken into the ledger accounts that sum
                            to it, with a TOTAL row that ties to the snapshot
  3. Ind AS Gaps          — standard | requirement | gap | INPUTS TAKEN | citations
  4. Ind AS Gap Accounts  — the accounts considered for each standard's assessment
  5. Findings             — ranked findings
  6. Quantified Risk Areas (only when document-sourced rows exist)

Styling mirrors trial_balance/tb_workbook.py so the two downloads look related.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from yukta_rag.audit.audit_config import MAX_GAP_ACCOUNTS_PER_SHEET
from yukta_rag.trial_balance.tb_tools import _net

_HEADER_FILL = PatternFill("solid", fgColor="1A3A5C")   # navy to match audit UI
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="DCE6F1")
_TOTAL_FILL = PatternFill("solid", fgColor="EDE9F6")
_HIGH_FILL = PatternFill("solid", fgColor="FADBD8")
_MED_FILL = PatternFill("solid", fgColor="FFF4CE")
_BOLD = Font(bold=True)
_NUM = "#,##0"
_THIN = Side(style="thin", color="C9D4E4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# standard number -> (FSLI lines, name keywords) whose accounts feed that gap's
# assessment; mirrors the trigger logic in audit_gaps.ind_as_gaps.
_GAP_ACCOUNT_SOURCES: dict[int, tuple[list[str], list[str]]] = {
    16: (["Property, Plant and Equipment"], []),
    23: (["Borrowings", "Capital Work-in-Progress"], []),
    2: (["Inventories"], []),
    115: (["Revenue from Operations"], []),
    109: (["Trade Receivables", "Loans (financial asset)", "Borrowings"], []),
    116: ([], ["lease", "right of use", "rou", "leasehold"]),
    36: (["Investments", "Property, Plant and Equipment", "Capital Work-in-Progress"], []),
    37: (["Provisions"], ["provision", "contingent", "write-off", "waiver"]),
    20: ([], ["grant", "subsidy"]),
    12: (["Deferred Tax Liability"], ["deferred tax"]),
    21: ([], ["foreign exchange", "forex", "exchange gain", "exchange loss",
              "fcnr", "ecb", "fctl"]),
    110: ([], ["subsidiary", "associate", "joint venture", "investment in "]),
    8: ([], ["suspense", "clearing", "prior period", "legacy", "migration", "adjustment"]),
}
_MAX_GAP_ACCOUNTS = MAX_GAP_ACCOUNTS_PER_SHEET


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _num_cols(ws, cols, first_row=2) -> None:
    for col in cols:
        for r in range(first_row, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = _NUM


def _fill_row(ws, ncols, fill, font=None) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.fill = fill
        if font:
            cell.font = font


def _fmt_inputs(inputs: dict) -> str:
    """Render a gap's trigger inputs for the Excel cell."""
    parts = []
    for k, v in (inputs or {}).items():
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            parts.append(f"{k}: {v:,.0f}")
        elif isinstance(v, bool):
            parts.append(f"{k}: {'yes' if v else 'no'}")
        elif isinstance(v, list):
            if v:
                parts.append(f"{k}: {', '.join(str(x) for x in v)}")
        else:
            parts.append(f"{k}: {v}")
    return "; ".join(parts)


def _sheet_snapshot(wb: Workbook, snapshot: dict, contributors: dict) -> None:
    ws = wb.active
    ws.title = "Financial Snapshot"
    periods = snapshot.get("periods", ["Balance"])
    headers = ["Metric"] + list(periods) + ["Contributing accounts"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for r in snapshot.get("rows", []):
        n_accts = len((contributors.get(r["metric"]) or {}).get("accounts", []))
        ws.append([r["metric"]] + list(r["values"]) + [n_accts])
    _num_cols(ws, range(2, 2 + len(periods)))
    _autosize(ws, {1: 32, **{1 + i: 20 for i in range(1, len(periods) + 1)},
                   2 + len(periods): 20})


def _sheet_drilldown(wb: Workbook, snapshot: dict, contributors: dict) -> None:
    ws = wb.create_sheet("Snapshot Drilldown")
    periods = snapshot.get("periods", ["Balance"])
    headers = ["Metric / Account", "Code", "FSLI", "Category"] + list(periods)
    ws.append(headers)
    _style_header(ws, len(headers))
    ncols = len(headers)
    for metric, data in contributors.items():
        ws.append([metric, "", "", ""] + list(data["values"]))
        _fill_row(ws, ncols, _SECTION_FILL, _BOLD)
        for acc in data["accounts"]:
            ws.append([f"    {acc['name']}", acc.get("code") or "",
                       acc.get("fsli") or "", acc.get("category") or ""]
                      + list(acc["contributions"]))
        ws.append([f"TOTAL — {metric}", "", "", ""] + list(data["values"]))
        _fill_row(ws, ncols, _TOTAL_FILL, _BOLD)
    _num_cols(ws, range(5, 5 + len(periods)))
    _autosize(ws, {1: 48, 2: 14, 3: 28, 4: 12,
                   **{4 + i: 20 for i in range(1, len(periods) + 1)}})


def _sheet_gaps(wb: Workbook, gaps: list[dict]) -> None:
    ws = wb.create_sheet("Ind AS Gaps")
    headers = ["Standard", "Name", "Requirement", "Gap indicated by the TB",
               "Inputs taken for the assessment", "Citations", "Evidence to request"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for g in gaps:
        cites = "; ".join(c.get("label", "") for c in g.get("citations", []))
        ws.append([g["standard"], g["name"], g["requirement"], g["gap"],
                   _fmt_inputs(g.get("inputs")), cites,
                   "; ".join(g.get("evidence_requested", []))])
    for r in range(2, ws.max_row + 1):
        for c in range(3, 8):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, {1: 12, 2: 26, 3: 40, 4: 48, 5: 44, 6: 40, 7: 34})


def _sheet_gap_accounts(wb: Workbook, gaps: list[dict], mapping: dict,
                        tb: dict) -> None:
    ws = wb.create_sheet("Ind AS Gap Accounts")
    periods = tb.get("periods", ["Balance"])
    headers = ["Standard", "Code", "Account", "FSLI"] + [f"Net ({p})" for p in periods]
    ws.append(headers)
    _style_header(ws, len(headers))
    ncols = len(headers)
    net_by_name = {a["name"]: a for a in tb["accounts"]}
    present_standards = []
    for g in gaps:
        try:
            present_standards.append(int(g["standard"].split()[-1]))
        except (ValueError, IndexError):
            continue
    for std in present_standards:
        fslis, kws = _GAP_ACCOUNT_SOURCES.get(std, ([], []))
        rows = []
        for a in mapping["accounts"]:
            name_l = a["name"].lower()
            if (a.get("fsli") in fslis) or (kws and any(k in name_l for k in kws)):
                src = net_by_name.get(a["name"])
                nets = [round(_net(src, i + 1), 2) if src else 0.0
                        for i in range(len(periods))]
                rows.append([a.get("code") or "", a["name"], a.get("fsli") or ""] + nets)
        rows.sort(key=lambda r: abs(r[3]), reverse=True)
        capped = len(rows) > _MAX_GAP_ACCOUNTS
        ws.append([f"Ind AS {std}", "", f"{len(rows)} account(s) considered"
                   + (f" — top {_MAX_GAP_ACCOUNTS} shown" if capped else ""), ""]
                  + [""] * len(periods))
        _fill_row(ws, ncols, _SECTION_FILL, _BOLD)
        for row in rows[:_MAX_GAP_ACCOUNTS]:
            ws.append([""] + row)
    _num_cols(ws, range(5, 5 + len(periods)))
    _autosize(ws, {1: 12, 2: 14, 3: 48, 4: 28,
                   **{4 + i: 20 for i in range(1, len(periods) + 1)}})


def _sheet_findings(wb: Workbook, findings: list[dict]) -> None:
    ws = wb.create_sheet("Findings")
    headers = ["#", "Rating", "Finding", "FSLI", "Amount", "Observation", "Gap",
               "Evidence requested"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for i, f in enumerate(findings, 1):
        ws.append([i, (f.get("risk_rating") or "").upper(), f.get("account", ""),
                   f.get("fsli") or "", f.get("amount"),
                   f.get("observation", ""), f.get("gap", ""),
                   "; ".join(f.get("evidence_requested", []))])
        rating = (f.get("risk_rating") or "").lower()
        if rating == "high":
            _fill_row(ws, len(headers), _HIGH_FILL)
        elif rating == "medium":
            _fill_row(ws, len(headers), _MED_FILL)
    _num_cols(ws, [5])
    for r in range(2, ws.max_row + 1):
        for c in (6, 7, 8):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, {1: 5, 2: 10, 3: 40, 4: 26, 5: 20, 6: 50, 7: 44, 8: 34})


def _sheet_quantified(wb: Workbook, quantified: dict) -> None:
    rows = (quantified or {}).get("rows") or []
    if not rows:
        return
    ws = wb.create_sheet("Quantified Risk Areas")
    headers = ["#", "Item", "Head affected", "Potential effect (per source)",
               "Amount", "Unit", "Source", "Verbatim quote"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for r in rows:
        src = r.get("source") or {}
        ws.append([r.get("n"), r.get("item", ""),
                   ", ".join(r.get("head_affected") or []),
                   r.get("effect_text", ""), r.get("amount"), r.get("unit") or "",
                   f"{src.get('filename', '')} p.{src.get('page_no', '')}",
                   r.get("quote", "")])
    cum = (quantified or {}).get("cumulative") or {}
    if cum.get("n_items"):
        per_dir = "; ".join(f"{k.replace('_', ' ')}: {v:,.2f}"
                            for k, v in (cum.get("by_direction") or {}).items())
        ws.append([])
        ws.append(["", f"CUMULATIVE ({cum['n_items']} item(s)"
                   + (f", {cum['unit']}" if cum.get("unit") else "") + ")",
                   "", per_dir, cum.get("gross_total"), cum.get("unit") or "",
                   "", cum.get("note", "")])
        _fill_row(ws, len(headers), _TOTAL_FILL, _BOLD)
    _num_cols(ws, [5])
    for r_i in range(2, ws.max_row + 1):
        for c in (2, 4, 8):
            ws.cell(row=r_i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, {1: 5, 2: 40, 3: 28, 4: 46, 5: 14, 6: 10, 7: 26, 8: 50})


def build_audit_workbook(entity: str, tb: dict, snapshot: dict, contributors: dict,
                         gaps: list[dict], mapping: dict, findings: list[dict],
                         quantified: dict | None = None) -> bytes:
    """Build TB_Audit.xlsx and return it as bytes (purely deterministic)."""
    wb = Workbook()
    _sheet_snapshot(wb, snapshot, contributors)
    _sheet_drilldown(wb, snapshot, contributors)
    _sheet_gaps(wb, gaps)
    _sheet_gap_accounts(wb, gaps, mapping, tb)
    _sheet_findings(wb, findings)
    _sheet_quantified(wb, quantified or {})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
