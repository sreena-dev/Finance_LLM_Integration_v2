"""Builds a {gl_code: hint_label} lookup from a company's own grouping
workbook. Grouping files range from a clean 2-column GL-code->group map to
hierarchical sheets where section/sub-section labels sit on separate
blank-code rows above the coded account rows, to a rich multi-level
hierarchy with its own structured columns.

Ported from TB_normalization_v1's input/grouping_source.py.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from modes.trial_balance.pipeline.tools.input_grouping_shape import GroupingShape, classify_grouping_shape
from modes.trial_balance.pipeline.tools.input_header_detect import locate_header, parse_gl_code
from modes.trial_balance.pipeline.tools.input_legacy_formats import load_grid_any_format
from modes.trial_balance.pipeline.tools.input_long_path import to_long_path

# Total/subtotal marker rows can be the bare word ("Total") or a compound
# label ("Total Note No. 2: Property plant and equipment", "SUB-TOTAL
# INTERUNIT"). Matching only the bare word misses the compound form
# entirely -- which both fails to pop the trail AND pushes the compound
# total text itself onto the trail as if it were a new section label.
_TOTAL_RE = re.compile(r"^(total\b|grand\s*total\b|sub[- ]?total\b)", re.IGNORECASE)
# Cross-reference/adjustment footnotes ("adj in GL 2051103") -- auditor
# annotations, not section headers, but they read exactly like one to the
# trail logic.
_ADJ_RE = re.compile(r"^adj\b", re.IGNORECASE)
_NOISE_LABELS = {"check", "control account", "-", "--"}
_MAX_TRAIL_DEPTH = 3  # a real section trail is rarely more than 2-3 levels deep


def _is_total_label(text: str) -> bool:
    return bool(_TOTAL_RE.match(text.strip()))


def _is_noise_label(text: str) -> bool:
    t = text.strip().lower()
    return t in _NOISE_LABELS or bool(_ADJ_RE.match(t))


def load_raw_grid(path, sheet_name: str = None) -> list:
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(to_long_path(path), read_only=True, data_only=True)
    except Exception:
        return load_grid_any_format(path, sheet_name=sheet_name)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return grid


def _is_numeric_like(value) -> bool:
    """True for real numbers and numeric-looking strings -- used to skip
    metadata columns like a "Note No." (2.1, 2.2, ...) that sit before the
    actual description column and would otherwise get picked as the
    label."""
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip()
    if not text:
        return True
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


STRUCTURED_FIELDS = ("bs_pl", "sub_head_2", "sub_head_1", "main_head", "account_type")

# A grouping cell can carry leftover note-reference numbers ("13.1"), a
# bullet-style prefix ("-  HDFC Bank"), or a leading list-marker number on
# otherwise-real text ("1-PROPERTY, PLANT AND EQUIPMENT"). None of these
# are real classification text on their own, but a bare .strip() lets all
# three through. Real line-item text always has letters and is more than a
# couple of characters, so blanking anything that isn't is safe.
_LEADING_DASH_RE = re.compile(r"^(?:\d+\s*[.\-]\s*|[\-–—\s])+")
_WHITESPACE_RUN_RE = re.compile(r"\s+")
_NUMERIC_ONLY_RE = re.compile(r"^-?\d+(\.\d+)?$")
_MIN_MEANINGFUL_LENGTH = 2


def normalize_grouping_cell_value(raw):
    """Cleans a raw grouping-cell value into display-ready text, or None
    if the cell is junk (a bare note-reference number, a leading bullet
    dash, or a fragment too short to be meaningful)."""
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    text = _LEADING_DASH_RE.sub("", text)
    text = _WHITESPACE_RUN_RE.sub(" ", text).strip()
    if not text or _NUMERIC_ONLY_RE.match(text) or len(text) < _MIN_MEANINGFUL_LENGTH:
        return None
    return text


def _looks_like_gl_code(code: str) -> bool:
    """True for a plausible GL code cell (purely numeric, or alphanumeric-
    no-space like "G00074"), false for a section-header row whose
    descriptive text happens to land in the code column instead of being
    blank. A real code is never a multi-word phrase."""
    return bool(code) and not any(ch.isspace() for ch in code)


def extract_structured_fields(grid: list, max_scan: int = 15) -> dict:
    """When the grouping file's own header already spells out any of the 5
    grouping columns, pull those values directly per coded GL row instead
    of relying on the LLM. Returns {} for grouping files with no such
    columns, which fall through to build_hint_map's text-hint path
    instead."""
    header_row, cols = locate_header(grid, max_scan)
    code_col = cols.get("gl_code")
    present_fields = [f for f in STRUCTURED_FIELDS if f in cols]
    if code_col is None or not present_fields:
        return {}

    structured: dict = {}
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        if code_col >= len(row):
            continue
        code = parse_gl_code(row[code_col])
        if not _looks_like_gl_code(code):
            continue
        fields = {}
        for f in present_fields:
            c = cols[f]
            if c < len(row) and row[c] not in (None, ""):
                cleaned = normalize_grouping_cell_value(row[c])
                if cleaned is not None:
                    fields[f] = cleaned
        if fields:
            structured[code] = fields

    return structured


def _hint_text_from_column(grid: list, header_row: int, code_col: int, source_col: int) -> dict:
    """Shared row-walk for both structured-taxonomy and flat-category
    shapes: pull one specific, already-identified column's text as the
    hint, keyed by GL code. No blind left-to-right scanning."""
    hints: dict = {}
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        code = parse_gl_code(row[code_col]) if code_col < len(row) else ""
        if not _looks_like_gl_code(code):
            continue
        if source_col < len(row) and row[source_col] not in (None, ""):
            hints[code] = str(row[source_col]).strip()
    return hints


def _hint_text_structured(grid: list, header_row: int, code_col: int, gl_name_col: int) -> dict:
    """STRUCTURED_TAXONOMY shape: hint text is the GL description column
    (already correctly identified by locate_header's alias table), not
    whatever the first non-numeric cell happens to be."""
    return _hint_text_from_column(grid, header_row, code_col, gl_name_col)


def _hint_text_flat_category(grid: list, header_row: int, code_col: int, group_label_col: int) -> dict:
    """FLAT_CATEGORY shape: hint text is the single flat category column
    (e.g. "Grouping")."""
    return _hint_text_from_column(grid, header_row, code_col, group_label_col)


def extract_hint_text(grid: list, max_scan: int = 15) -> dict:
    """Single entry point for hint_text extraction. Classifies the
    grouping file's shape once and dispatches to exactly one strategy --
    never unconditionally runs build_hint_map on every file regardless of
    shape, which would silently degrade shapes it wasn't designed for."""
    header_row, cols = locate_header(grid, max_scan)
    code_col = cols.get("gl_code")
    if code_col is None:
        return {}

    shape = classify_grouping_shape(cols)
    if shape is GroupingShape.STRUCTURED_TAXONOMY and "gl_name" in cols:
        return _hint_text_structured(grid, header_row, code_col, cols["gl_name"])
    if shape is GroupingShape.FLAT_CATEGORY:
        return _hint_text_flat_category(grid, header_row, code_col, cols["group_label"])
    # HIERARCHICAL_TRAIL, or STRUCTURED_TAXONOMY with taxonomy columns but
    # no gl_name/description column at all -- both fall back to the trail
    # scan rather than returning {} and losing all signal.
    return build_hint_map(grid, max_scan)


def build_hint_map(grid: list, max_scan: int = 15) -> dict:
    header_row, cols = locate_header(grid, max_scan)
    code_col = cols.get("gl_code")
    if code_col is None:
        return {}

    hints: dict = {}
    trail: list = []

    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        code_val = row[code_col] if code_col < len(row) else None
        code = parse_gl_code(code_val)
        is_coded = _looks_like_gl_code(code)

        text_cells = [
            str(v).strip()
            for c, v in enumerate(row)
            if c != code_col and v not in (None, "") and not _is_numeric_like(v)
        ]
        text_cells = [t for t in text_cells if not _is_noise_label(t)]
        is_total_row = any(_is_total_label(t) for t in text_cells)

        if is_coded:
            label = text_cells[0] if text_cells else ""
            parts = [*trail, label] if label else list(trail)
            if parts:
                hints[code] = " / ".join(dict.fromkeys(parts))
        elif is_total_row:
            # a Total row closes out the section opened by the most
            # recently pushed label, so pop it rather than letting it
            # linger and bleed into the next sibling section's hints
            if trail:
                trail.pop()
        elif text_cells:
            trail.append(text_cells[0])
            if len(trail) > _MAX_TRAIL_DEPTH:
                trail.pop(0)

    return hints
