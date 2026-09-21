"""Format-agnostic workbook loading: `.xlsx`/`.xlsm` (openpyxl), genuine
legacy `.xls` (xlrd), and `.xlsb` (pyxlsb).

Content-sniffed, NEVER extension-based -- a real-world export's extension
can lie (a file named `..._Trial_Balance.xls` can, by its actual bytes, be
a perfectly normal `.xlsx` zip that was just given the wrong extension;
trusting the extension there would route it into xlrd, which correctly
refuses real xlsx content). Only a genuine OLE2/BIFF file (the old
pre-2007 `.xls` binary format, magic bytes D0 CF 11 E0 A1 B1 1A E1) needs
xlrd; a zip archive containing `xl/workbook.bin` (binary records, not XML)
is `.xlsb` and needs pyxlsb; a zip archive containing `xl/workbook.xml` is
real `.xlsx`.

Ported from TB_normalization_v1's input/legacy_formats.py.
"""

from __future__ import annotations

import io
import zipfile
from enum import Enum
from pathlib import Path
from typing import Optional

from modes.trial_balance.pipeline.tools.input_long_path import to_long_path

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class UnsupportedWorkbookFormatError(Exception):
    """The workbook's actual content doesn't match any format this module
    can read, or the optional dependency needed to read it (xlrd/pyxlsb)
    isn't installed."""


class PasswordProtectedError(Exception):
    """The workbook is a genuinely encrypted OOXML/OLE2 file (not a plain
    legacy .xls BIFF file) -- reading it requires the password, which this
    module never has. Distinguished from a real legacy .xls by content
    (msoffcrypto's own EncryptionInfo-stream check), not the OLE2 magic
    bytes alone: an encrypted .xlsx is ALSO OLE2-wrapped (that's how Office
    stores an encrypted container), so _sniff_format's magic-byte check on
    its own cannot tell the two apart."""


class _WorkbookFormat(Enum):
    XLSX = "xlsx"       # openpyxl -- ZIP archive containing xl/workbook.xml
    XLSB = "xlsb"        # pyxlsb  -- ZIP archive containing xl/workbook.bin
    LEGACY_XLS = "xls"   # xlrd    -- OLE2/BIFF binary, not a ZIP at all


def _is_encrypted_ole2(path: Path) -> bool:
    """True iff an OLE2-magic file is actually an encrypted OOXML container
    (has an EncryptionInfo stream), not a genuine legacy .xls. Best-effort:
    msoffcrypto not being installed, or any read failure, means "assume not
    encrypted" -- falls through to the existing xlrd path exactly as before,
    never a new failure mode of its own."""
    try:
        import msoffcrypto

        with open(path, "rb") as f:
            return msoffcrypto.OfficeFile(f).is_encrypted()
    except ImportError:
        return False
    except Exception:
        return False


def _sniff_format(path: Path) -> _WorkbookFormat:
    with open(path, "rb") as f:
        head = f.read(8)
    if head == _OLE2_MAGIC:
        if _is_encrypted_ole2(path):
            raise PasswordProtectedError(
                f"{path.name} is password-protected. Please remove the password and re-upload."
            )
        return _WorkbookFormat.LEGACY_XLS
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
        if "xl/workbook.bin" in names:
            return _WorkbookFormat.XLSB
        if "xl/workbook.xml" in names:
            return _WorkbookFormat.XLSX
    raise UnsupportedWorkbookFormatError(
        f"{path.name}: unrecognized workbook format (not OLE2/.xls, not a "
        f".xlsx/.xlsm zip, not a .xlsb zip)"
    )


def _load_via_openpyxl(path: Path, sheet_name: Optional[str], max_row: Optional[int]) -> dict:
    import openpyxl

    with open(path, "rb") as f:
        data = f.read()
    # BytesIO, not the path -- openpyxl's own format validation keys off the
    # file extension when given a path, which is exactly the signal already
    # proven unreliable here. A stream sidesteps that check entirely; the
    # content is already confirmed to be a real xlsx zip by _sniff_format.
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = [sheet_name] if sheet_name else wb.sheetnames
        return {name: [list(row) for row in wb[name].iter_rows(values_only=True, max_row=max_row)] for name in names}
    finally:
        wb.close()


def _load_via_xlrd(path: Path, sheet_name: Optional[str], max_row: Optional[int]) -> dict:
    try:
        import xlrd
    except ImportError as e:
        raise UnsupportedWorkbookFormatError(
            f"{path.name} is a legacy .xls (BIFF) file -- install xlrd to read it."
        ) from e

    wb = xlrd.open_workbook(str(path))
    names = [sheet_name] if sheet_name else wb.sheet_names()
    result: dict = {}
    for name in names:
        sheet = wb.sheet_by_name(name)
        n_rows = sheet.nrows if max_row is None else min(sheet.nrows, max_row)
        grid: list = []
        for r in range(n_rows):
            row_cells = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                row_cells.append(_normalize_xlrd_cell(cell, wb.datemode))
            grid.append(row_cells)
        result[name] = grid
    return result


def _normalize_xlrd_cell(cell, datemode: int):
    import xlrd

    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    return cell.value


def _load_via_pyxlsb(path: Path, sheet_name: Optional[str], max_row: Optional[int]) -> dict:
    try:
        import pyxlsb
    except ImportError as e:
        raise UnsupportedWorkbookFormatError(
            f"{path.name} is a .xlsb (Excel Binary Workbook) file -- install pyxlsb to read it."
        ) from e

    result: dict = {}
    with pyxlsb.open_workbook(str(path)) as wb:
        names = [sheet_name] if sheet_name else wb.sheets
        for name in names:
            with wb.get_sheet(name) as sheet:
                grid: list = []
                for row in sheet.rows():
                    if max_row is not None and len(grid) >= max_row:
                        break
                    if not row:
                        grid.append([])
                        continue
                    width = max(c.c for c in row) + 1
                    row_values: list = [None] * width
                    for c in row:
                        row_values[c.c] = c.v
                    grid.append(row_values)
            result[name] = grid
    return result


def load_all_sheets_any_format(
    path: Path, *, sheet_name: Optional[str] = None, max_row: Optional[int] = None
) -> dict:
    """Content-sniffs `path` and returns {sheet_name: grid} for every sheet
    (or just `sheet_name` if given), regardless of actual format --
    .xlsx/.xlsm via openpyxl, .xlsb via pyxlsb, genuine legacy .xls via
    xlrd. Every grid is the same row-major list[list] shape, with blank
    cells normalized to None across all three backends."""
    resolved = to_long_path(path)
    fmt = _sniff_format(resolved)
    if fmt is _WorkbookFormat.XLSX:
        return _load_via_openpyxl(resolved, sheet_name, max_row)
    if fmt is _WorkbookFormat.XLSB:
        return _load_via_pyxlsb(resolved, sheet_name, max_row)
    return _load_via_xlrd(resolved, sheet_name, max_row)


def load_grid_any_format(path: Path, *, sheet_name: Optional[str] = None, max_row: Optional[int] = None) -> list:
    """Single-grid convenience wrapper: returns the first sheet's grid (or
    `sheet_name`'s, if given) instead of the full {name: grid} dict."""
    sheets = load_all_sheets_any_format(path, sheet_name=sheet_name, max_row=max_row)
    if sheet_name is not None:
        return sheets[sheet_name]
    return next(iter(sheets.values()))
