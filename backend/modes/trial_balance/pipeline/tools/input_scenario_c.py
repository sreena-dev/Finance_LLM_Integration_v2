"""Scenario C adapter: a single workbook where TB rows and grouping data
live in the SAME sheet/grid -- e.g. one flat "Grouping" category column
sitting alongside the TB's own gl_code/gl_name/opening/debit/credit/
closing columns (one sheet, no separate grouping file, no COMPANY_DETAILS/
TB&GROUPING-TEMPLATE sheets either).

Distinct from Scenario A (a specific clean structured-taxonomy template)
and Scenario B (two separate files) -- this is what a caller-supplied
single file that isn't shaped like Scenario A's template needs.

Reuses input_grouping_source.py's extract_structured_fields/
extract_hint_text unchanged -- the combined-file case is a ROUTING
difference (one grid serving both the TB-row role and the grouping role,
instead of two grids), not a new content shape.

Ported from TB_normalization_v1's input/tb_grouping_combined.py.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import Optional

import openpyxl

from modes.trial_balance.pipeline.tools.document_metadata import DocumentMetadata, blank_metadata, with_resolved_framework
from modes.trial_balance.pipeline.tools.input_grouping_shape import classify_grouping_shape
from modes.trial_balance.pipeline.tools.input_grouping_source import extract_hint_text, extract_structured_fields
from modes.trial_balance.pipeline.tools.input_header_detect import build_tb_rows_from_grid, pick_tb_sheet
from modes.trial_balance.pipeline.tools.input_legacy_formats import load_all_sheets_any_format
from modes.trial_balance.pipeline.tools.input_long_path import to_long_path
from modes.trial_balance.pipeline.tools.parsed_input import ParsedInput
from modes.trial_balance.pipeline.tools.tb_models import GroupingHint


class CombinedWorkbookStructureError(Exception):
    """The single supplied workbook has no TB-qualifying sheet (no
    gl_code + closing column found on any sheet)."""


def parse_combined_workbook(path: Path, *, metadata: Optional[DocumentMetadata] = None) -> ParsedInput:
    """`metadata` is caller-supplied, same as parse_tb_grouping_pair --
    there is no COMPANY_DETAILS sheet in this scenario either."""
    try:
        wb = openpyxl.load_workbook(to_long_path(path), read_only=True, data_only=True)
    except Exception:
        sheets = load_all_sheets_any_format(path)
    else:
        try:
            sheets = {name: [list(row) for row in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}
        finally:
            wb.close()

    pick = pick_tb_sheet(sheets)
    if pick is None:
        raise CombinedWorkbookStructureError(f"{path.name}: no TB-qualifying sheet found")
    sheet_name, header_row, cols = pick
    grid = sheets[sheet_name]

    tb_rows = build_tb_rows_from_grid(grid, header_row, cols)

    # Same grid, second pass: the shape-aware grouping extraction works
    # unchanged here -- classify_grouping_shape sees the same taxonomy/
    # category/neither columns it would see reading a standalone grouping
    # file.
    structured = extract_structured_fields(grid)
    hint_text_map = extract_hint_text(grid)
    all_codes = set(structured) | set(hint_text_map)
    grouping_hints = {
        code: GroupingHint(gl_code=code, known_fields=structured.get(code, {}), hint_text=hint_text_map.get(code, ""))
        for code in all_codes
    }

    if metadata is None:
        metadata = blank_metadata()
    metadata = with_resolved_framework(metadata, explicit=metadata.framework, heuristic_grid=grid)
    metadata = replace(metadata, tb_doc_name=path.name, grouping_doc_name=path.name, has_grouping=bool(grouping_hints))

    return ParsedInput(
        metadata=metadata, tb_rows=tb_rows, grouping_hints=grouping_hints, grouping_shape=classify_grouping_shape(cols),
    )
