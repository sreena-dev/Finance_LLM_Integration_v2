"""Scenario B adapter: two separate uploaded workbooks (a TB file and a
Grouping file) -> ParsedInput.

Workbook identification is content-based, never filename-based -- reuses
input_header_detect.pick_tb_sheet (the same signal Scenario A's header
detection is built on) as the sole determination: whichever file has a
sheet with a gl_code + closing column is the TB file; the other is the
grouping file.

Ported from TB_normalization_v1's input/tb_grouping_pair.py.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import Optional

import openpyxl

from modes.trial_balance.pipeline.tools.document_metadata import DocumentMetadata, blank_metadata, with_resolved_framework
from modes.trial_balance.pipeline.tools.input_grouping_shape import classify_grouping_shape
from modes.trial_balance.pipeline.tools.input_grouping_source import extract_hint_text, extract_structured_fields, load_raw_grid
from modes.trial_balance.pipeline.tools.input_header_detect import build_tb_rows_from_grid, locate_header, pick_tb_sheet
from modes.trial_balance.pipeline.tools.input_legacy_formats import load_all_sheets_any_format
from modes.trial_balance.pipeline.tools.input_long_path import to_long_path
from modes.trial_balance.pipeline.tools.parsed_input import ParsedInput
from modes.trial_balance.pipeline.tools.tb_models import GroupingHint, TBRow  # noqa: F401 -- TBRow used by type hints/callers


class InputClassificationError(Exception):
    """Content-based TB-vs-grouping workbook classification did not
    resolve to exactly one of each, or the identified TB workbook has no
    TB-qualifying sheet on a full parse."""


class DuplicateTbNoGroupingError(InputClassificationError):
    """>=2 files independently qualify as TB-shaped, no non-TB file exists
    among the inputs, but the TB candidates all parse to identical rows (a
    genuine duplicate export, not an ambiguity) -- `resolved_tb_path` is
    the single canonical file to use. There is still no grouping file
    among what was given, though: this is not, by itself, a resolvable
    Scenario B pair."""

    def __init__(self, resolved_tb_path: Path, duplicate_paths: list):
        self.resolved_tb_path = resolved_tb_path
        self.duplicate_paths = duplicate_paths
        super().__init__(
            f"{len(duplicate_paths)} files are duplicate TB exports (identical parsed rows) with no "
            f"grouping file among them -- resolved to {resolved_tb_path.name}, but no grouping source "
            f"was found in this input set: {[p.name for p in duplicate_paths]}"
        )


def _load_workbook_sheets(path: Path, max_row: Optional[int] = None) -> dict:
    """openpyxl direct-path load is the default -- unchanged behavior and
    performance for every genuine .xlsx/.xlsm file. Falls back to
    legacy_formats' content-sniffing loader only when the direct load
    fails."""
    try:
        wb = openpyxl.load_workbook(to_long_path(path), read_only=True, data_only=True)
    except Exception:
        return load_all_sheets_any_format(path, max_row=max_row)
    try:
        return {name: [list(row) for row in wb[name].iter_rows(values_only=True, max_row=max_row)] for name in wb.sheetnames}
    finally:
        wb.close()


def _select_primary_tb_candidate(tb_candidates: list) -> Path:
    """Only called when >=2 files independently qualify as TB-shaped and
    no non-TB file exists among the inputs. The one signal that's actually
    safe is content equality: if every candidate parses to the exact same
    TB rows, picking any one of them is not a guess. If the parsed rows
    genuinely differ, there's no content-based way to know which is
    authoritative -- refuse rather than guess."""
    first_rows = _read_tb_rows(tb_candidates[0])
    for other in tb_candidates[1:]:
        if _read_tb_rows(other) != first_rows:
            raise InputClassificationError(
                f"{len(tb_candidates)} files independently qualify as TB-shaped and no "
                f"non-TB file exists among the inputs, but their parsed rows differ -- "
                f"refusing to guess which is authoritative: {[f.name for f in tb_candidates]}"
            )
    # All candidates parse identically -- deterministic (not arbitrary)
    # pick: shortest filename.
    return min(tb_candidates, key=lambda p: (len(p.name), p.name))


def classify_workbooks(files: list) -> tuple:
    """(tb_path, grouping_path). Bounded scan (max_row=30) for speed --
    classification only needs to see the header + a handful of data rows,
    not the whole workbook."""
    tb_candidates: list = []
    other: list = []
    for f in files:
        sheets = _load_workbook_sheets(f, max_row=30)
        if pick_tb_sheet(sheets):
            tb_candidates.append(f)
        else:
            other.append(f)

    if len(tb_candidates) > 1 and not other:
        resolved = _select_primary_tb_candidate(tb_candidates)
        raise DuplicateTbNoGroupingError(resolved, tb_candidates)

    if len(tb_candidates) != 1 or len(other) != 1:
        raise InputClassificationError(
            f"Expected exactly one TB workbook and one grouping workbook among "
            f"{[f.name for f in files]}; got TB candidates={[f.name for f in tb_candidates]}, "
            f"other={[f.name for f in other]}"
        )
    return tb_candidates[0], other[0]


def _read_tb_rows(path: Path) -> list:
    sheets = _load_workbook_sheets(path)  # full parse this time, no max_row cap
    pick = pick_tb_sheet(sheets)
    if pick is None:
        raise InputClassificationError(f"{path.name}: no TB-qualifying sheet found on full parse")
    sheet_name, header_row, cols = pick
    return build_tb_rows_from_grid(sheets[sheet_name], header_row, cols)


def read_grouping_hints(grouping_grid: list) -> dict:
    """known_fields (extract_structured_fields) stays unconditional --
    it's already column-name-aware and correct for every shape, returning
    {} harmlessly for files with no discrete taxonomy columns. hint_text
    (extract_hint_text) is shape-dispatched: exactly one strategy per
    file, not a merge of every extractor's output."""
    structured = extract_structured_fields(grouping_grid)
    hint_text_map = extract_hint_text(grouping_grid)
    all_codes = set(structured) | set(hint_text_map)
    return {
        code: GroupingHint(gl_code=code, known_fields=structured.get(code, {}), hint_text=hint_text_map.get(code, ""))
        for code in all_codes
    }


def parse_tb_grouping_pair(
    tb_path: Path, grouping_path: Path, *, metadata: Optional[DocumentMetadata] = None,
) -> ParsedInput:
    """`metadata` is caller-supplied -- there is no COMPANY_DETAILS sheet
    in this scenario, so entity/company/FY fields must come from the
    caller. Framework is always (re-)resolved here against the grouping
    file's own vocabulary as the heuristic fallback, same precedence rule
    as the template path."""
    tb_rows = _read_tb_rows(tb_path)
    grouping_grid = load_raw_grid(grouping_path)
    grouping_hints = read_grouping_hints(grouping_grid)
    _, grouping_cols = locate_header(grouping_grid)
    grouping_shape = classify_grouping_shape(grouping_cols)

    if metadata is None:
        metadata = blank_metadata()

    metadata = with_resolved_framework(metadata, explicit=metadata.framework, heuristic_grid=grouping_grid)
    metadata = replace(metadata, tb_doc_name=tb_path.name, grouping_doc_name=grouping_path.name, has_grouping=True)

    return ParsedInput(metadata=metadata, tb_rows=tb_rows, grouping_hints=grouping_hints, grouping_shape=grouping_shape)
