"""Wave 9 remark #16: a joined "GL1; GL2; GL3; ..." requested-for cell is hard to
work row-by-row when a finding legitimately spans many accounts -- cap the display
instead of running every account together in one unreadable string."""

import json

from openpyxl import Workbook

from modes.trial_balance.pipeline.tools.reports import _write_findings_and_evidence_sheets


def _load_json_factory(out_dir):
    def load_json(explicit_path, default_filename):
        path = explicit_path or (out_dir / default_filename)
        if not path.exists():
            return {}
        return json.loads(path.read_text(encoding="utf-8"))
    return load_json


def test_requested_for_caps_long_account_lists(tmp_path):
    (tmp_path / "finding_records.json").write_text(json.dumps({
        "finding_records": [
            {"account": "GL1 - Test", "fsli": "Test", "risk_rating": "high",
             "observation": "o", "expectation": "e", "gap": "g",
             "evidence_requested": ["Bank statement"]},
        ],
    }), encoding="utf-8")
    (tmp_path / "evidence_request_list.json").write_text(json.dumps({
        "evidence_request_list": [
            {"record": "Bank statement", "audit_area": "Cash & Bank", "highest_risk_rating": "high",
             "regularity_relevant": False,
             "requested_for": ["GL1", "GL2", "GL3", "GL4", "GL5"],
             "assertions_resolved": ["Existence"]},
        ],
        "information_request_list": [],
    }), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)

    def new_sheet(name, row_count=0):
        return wb.create_sheet(name)

    _write_findings_and_evidence_sheets(new_sheet, tmp_path, _load_json_factory(tmp_path))

    ws = wb["Evidence Request Register"]
    requested_for_cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("GL1")]
    assert requested_for_cells, "expected a rendered 'Requested For' cell starting with GL1"
    cell_value = requested_for_cells[0]
    assert cell_value == "GL1; GL2; GL3; +2 more"


def test_requested_for_shows_all_accounts_when_within_cap(tmp_path):
    (tmp_path / "finding_records.json").write_text(json.dumps({
        "finding_records": [
            {"account": "GL1 - Test", "fsli": "Test", "risk_rating": "high",
             "observation": "o", "expectation": "e", "gap": "g",
             "evidence_requested": ["Bank statement"]},
        ],
    }), encoding="utf-8")
    (tmp_path / "evidence_request_list.json").write_text(json.dumps({
        "evidence_request_list": [
            {"record": "Bank statement", "audit_area": "Cash & Bank", "highest_risk_rating": "high",
             "regularity_relevant": False,
             "requested_for": ["GL1", "GL2"],
             "assertions_resolved": ["Existence"]},
        ],
        "information_request_list": [],
    }), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)

    def new_sheet(name, row_count=0):
        return wb.create_sheet(name)

    _write_findings_and_evidence_sheets(new_sheet, tmp_path, _load_json_factory(tmp_path))

    ws = wb["Evidence Request Register"]
    requested_for_cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("GL1")]
    assert requested_for_cells[0] == "GL1; GL2"
