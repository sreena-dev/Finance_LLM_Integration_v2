"""Round-trips custom_field_1/2/3 (reserved/not-yet-defined columns that
already exist on tb_table/live_tb_table) through ingest_tb_to_live's
custom_fields parameter, then a LIVE read-back -- confirming they survive
un-mutated when supplied and stay None when absent.

Rewritten from the retired build_canonical_tb/persist_canonical_tb_to_live
pair: ingest_tb_to_live is now the single entry point for both live and
uploaded documents (see backend/tools/db_bridge.py), so custom_fields moved
there as a first-class parameter instead."""

from unittest.mock import patch

import openpyxl
import polars as pl
import pytest

from modes.trial_balance.pipeline.tools import delete_db_document, ingest_tb_to_live


class _FakeLLMClient:
    def generate(self, *args, **kwargs):
        raise AssertionError("LLM should not be called for a fully-mapped template")


class _FakeAgent:
    llm_client = _FakeLLMClient()


def _write_template(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    company = wb.create_sheet("COMPANY_DETAILS")
    for row in [
        ["Company Details", None],
        ["COMPANY_NAME", "Custom Field Test Co"],
        ["FINANCIAL_YEAR", "FY2025-26"],
        ["FY_START", 2025],
        ["FY_END", 2026],
    ]:
        company.append(row)

    tb = wb.create_sheet("TB&GROUPING-TEMPLATE")
    tb.append(["GL Code", "GL Name", "Opening Balance", "Debit", "Credit", "Closing Balance",
               "BS/PL", "Main Head", "Sub Head 1", "Sub Head 2"])
    tb.append(["CF-1001", "Freehold Land (custom-field test)", 0, 1000, 0, 1000,
               "BS", "Non-current assets", "Property, Plant and Equipment", "Land"])
    wb.save(path)


def test_ingest_tb_to_live_populates_supplied_custom_fields(db_available, tmp_path):
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
            custom_fields={"custom_field_1": "custom-value-1", "custom_field_2": "custom-value-2"},
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        df = pl.read_parquet(tmp_path / "out" / "canonical_tb.parquet")
        assert set(df["custom_field_1"].to_list()) == {"custom-value-1"}
        assert set(df["custom_field_2"].to_list()) == {"custom-value-2"}
        assert set(df["custom_field_3"].to_list()) == {None}
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])


def test_ingest_tb_to_live_leaves_custom_fields_null_when_not_supplied(db_available, tmp_path):
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(tb_grouping_template_path=str(path), output_dir=str(tmp_path / "out"))

    try:
        assert result["execution_status"] == "SUCCESS", result
        df = pl.read_parquet(tmp_path / "out" / "canonical_tb.parquet")
        for col in ("custom_field_1", "custom_field_2", "custom_field_3"):
            assert set(df[col].to_list()) == {None}
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])


def test_custom_fields_survive_end_to_end_into_live_db(db_available, tmp_path):
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
            custom_fields={"custom_field_1": "e2e-value-1"},
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        tb_doc_id = result["tb_doc_id"]

        from modes.trial_balance.pipeline.db import db_cursor

        with db_cursor() as cur:
            cur.execute(
                "SELECT gl_code, custom_field_1 FROM live_tb_table WHERE tb_doc_id = %s ORDER BY gl_code",
                (tb_doc_id,),
            )
            rows = [dict(r) for r in cur.fetchall()]

        assert len(rows) == 1
        assert rows[0]["custom_field_1"] == "e2e-value-1"
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])
