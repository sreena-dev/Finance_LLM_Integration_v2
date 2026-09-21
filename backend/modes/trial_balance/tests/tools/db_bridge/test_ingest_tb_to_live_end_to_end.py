"""End-to-end integration test for the rewired ingest_tb_to_live: real
synthetic workbook -> input_dispatch auto-detection -> the native
classification engine (against the real shared taxonomy DB) -> a real
write into live_document_table/live_tb_table. Gated on `db_available`,
same as every other LIVE-staging test in this package.
"""

from unittest.mock import patch

import openpyxl
import pytest

from modes.trial_balance.pipeline.db import fetch_live_document
from modes.trial_balance.pipeline.tools import delete_db_document, ingest_tb_to_live


class _FakeResponse:
    def __init__(self, content):
        self.content = content


class _FakeLLMClient:
    """Only Step 1 (no hint) and Step 2 (complete, valid hint) rows are
    exercised by this fixture's data, so the LLM should never actually be
    called -- if it is, something upstream regressed."""

    def generate(self, *args, **kwargs):
        raise AssertionError("LLM should not be called for a fully-mapped template")


class _FakeAgent:
    llm_client = _FakeLLMClient()


def _write_template(path, tb_doc_name="ACME_FY2025-26.xlsx", with_grand_total_row=False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    company = wb.create_sheet("COMPANY_DETAILS")
    for row in [
        ["Company Details", None],
        ["COMPANY_NAME", "ACME End To End Ltd"],
        ["FINANCIAL_YEAR", "FY2025-26"],
        ["FY_START", 2025],
        ["FY_END", 2026],
    ]:
        company.append(row)

    tb = wb.create_sheet("TB&GROUPING-TEMPLATE")
    tb.append(["GL Code", "GL Name", "Opening Balance", "Debit", "Credit", "Closing Balance",
               "BS/PL", "Main Head", "Sub Head 1", "Sub Head 2"])
    tb.append(["E2E-1001", "Freehold Land (E2E test)", 0, 1000, 0, 1000,
               "BS", "Non-current assets", "Property, Plant and Equipment", "Land"])
    if with_grand_total_row:
        # Gap 1: a Grand Total row in the TB workbook itself must never reach
        # canonical_tb.parquet as a real GL account.
        tb.append(["E2E-9999", "Grand Total", 0, 1000, 0, 1000, None, None, None, None])
    wb.save(path)


def test_template_ingests_end_to_end_into_live_staging(db_available, tmp_path):
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        assert result["pipeline_status"] in ("SUCCESS", "WARNING"), result
        assert result["quality_tier"] == "CASE_1"
        assert result["run_metrics"]["mapped_count"] == 1
        assert (tmp_path / "out" / "canonical_tb.parquet").exists()

        import polars as pl

        out_df = pl.read_parquet(tmp_path / "out" / "canonical_tb.parquet")
        row = out_df.to_dicts()[0]
        assert row["gl_code"] == "E2E-1001"
        assert row["mapped_status"] == "MAPPED"
        assert row["sub_head_2"] == "Land"
        assert row["account_type"] == "Asset"

        # Blank overrides (the default -- no Company Details fields filled) leave
        # the workbook's own COMPANY_DETAILS metadata exactly as auto-parsed.
        live_doc = fetch_live_document(result["tb_doc_id"])
        assert live_doc["company_name"] == "ACME End To End Ltd"
        assert live_doc["financial_year"] == "FY2025-26"
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])


def test_company_details_override_wins_over_parsed_metadata(db_available, tmp_path):
    """A user-supplied Company Details field (company_name/cin/financial_year)
    must win over the workbook's own COMPANY_DETAILS sheet -- and the generated
    tb_doc_id must reflect the correction too, not just the display fields,
    since it's derived from company_name/financial_year (see build_tb_doc_id)."""
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
            company_name="Override Corp Ltd",
            cin="U00000MH2000GOI000001",
            financial_year="2030-31",
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        live_doc = fetch_live_document(result["tb_doc_id"])
        assert live_doc["company_name"] == "Override Corp Ltd"
        assert live_doc["cin"] == "U00000MH2000GOI000001"
        assert live_doc["financial_year"] == "2030-31"
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])


def test_company_details_override_applies_per_field_independently(db_available, tmp_path):
    """Only company_name is overridden here -- cin/financial_year must stay
    exactly as the workbook's own COMPANY_DETAILS sheet produced them."""
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
            company_name="Override Corp Ltd",
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        live_doc = fetch_live_document(result["tb_doc_id"])
        assert live_doc["company_name"] == "Override Corp Ltd"
        assert live_doc["financial_year"] == "FY2025-26"  # unchanged, no override given
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])


def test_persist_to_live_false_never_writes_to_live(db_available, tmp_path):
    """The query-analysis staging path (a future feature builds on this): the
    exact same classify/quality-gate chain runs and canonical_tb.parquet is
    still written, but nothing reaches live_document_table/live_tb_table --
    this is the direct behavioral contract that feature depends on."""
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
            persist_to_live=False,
        )

    assert result["execution_status"] == "SUCCESS", result
    assert result["persisted_to_live"] is False
    assert result["db_refs"] is None
    assert result["replaced_existing_document"] is None
    assert (tmp_path / "out" / "canonical_tb.parquet").exists()

    # The defining assertion -- nothing written to LIVE for this tb_doc_id.
    assert fetch_live_document(result["tb_doc_id"]) is None


def test_grand_total_row_never_reaches_canonical_tb(db_available, tmp_path):
    if not db_available:
        pytest.skip("No reachable Postgres DB (settings.DB_*) -- skipping.")

    path = tmp_path / "template.xlsx"
    _write_template(path, with_grand_total_row=True)

    with patch("modes.trial_balance.pipeline.agent.get_agent", return_value=_FakeAgent()):
        result = ingest_tb_to_live(
            tb_grouping_template_path=str(path),
            output_dir=str(tmp_path / "out"),
        )

    try:
        assert result["execution_status"] == "SUCCESS", result
        # Still exactly 1 real GL row -- the Grand Total row was excluded at
        # parse time, before the engine (and CASE_1's zero-LLM-touch routing)
        # ever saw it, so this stays CASE_1 rather than falling to CASE_2/3.
        assert result["quality_tier"] == "CASE_1"
        assert result["run_metrics"]["mapped_count"] == 1

        import polars as pl

        out_df = pl.read_parquet(tmp_path / "out" / "canonical_tb.parquet")
        assert out_df.height == 1
        assert out_df.to_dicts()[0]["gl_code"] == "E2E-1001"
    finally:
        delete_db_document(tb_doc_id=result["tb_doc_id"])
