"""Wave 9 remark #21 (Part E): the Ratio Trend sheet renders PY value / CY value /
direction per ratio -- a render job over both periods' financial_ratios.json, no new
ratio computation."""

import json

from openpyxl import Workbook

from modes.trial_balance.pipeline.tools.comparison import _ratio_trend_sheet


def test_direction_computed_from_both_periods(tmp_path):
    py_dir = tmp_path / "py"
    cy_dir = tmp_path / "cy"
    py_dir.mkdir()
    cy_dir.mkdir()
    (py_dir / "financial_ratios.json").write_text(json.dumps({
        "ratios": {"current": {"value": 1.2}, "debt_equity": {"value": 0.8}},
    }), encoding="utf-8")
    (cy_dir / "financial_ratios.json").write_text(json.dumps({
        "ratios": {"current": {"value": 1.5}, "debt_equity": {"value": 0.5}},
    }), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)

    def new_sheet(name, row_count=0):
        return wb.create_sheet(name)

    _ratio_trend_sheet(new_sheet, py_dir, cy_dir, "FY2024-25", "FY2025-26")

    ws = wb["Ratio Trend (PY vs CY)"]
    rows = {row[0].value: (row[1].value, row[2].value, row[3].value) for row in ws.iter_rows(min_row=5) if row[0].value}
    assert rows["Current"] == ("1.20", "1.50", "Up")
    assert rows["Debt Equity"] == ("0.80", "0.50", "Down")


def test_no_sheet_written_when_neither_period_has_ratios(tmp_path):
    py_dir = tmp_path / "py"
    cy_dir = tmp_path / "cy"
    py_dir.mkdir()
    cy_dir.mkdir()

    wb = Workbook()
    wb.remove(wb.active)

    def new_sheet(name, row_count=0):
        return wb.create_sheet(name)

    _ratio_trend_sheet(new_sheet, py_dir, cy_dir, "FY2024-25", "FY2025-26")

    assert "Ratio Trend (PY vs CY)" not in wb.sheetnames
