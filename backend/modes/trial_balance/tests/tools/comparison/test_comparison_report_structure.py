"""The comparative report: the agreed 25-section structure and the cross-reference
invariant, asserted over a real PY/CY run.

The comparative writer builds its Word document BEFORE its workbook, so the schedule
manifest does not exist while sections 1-11 render. The document save is therefore
deferred until after the workbook is written, and sections 12-25 plus the schedule
appendix render at that point. These tests exist because that ordering is easy to
break silently -- a re-ordered save would produce a document whose appendix is empty
while the workbook is full, and nothing would raise.
"""

import json
import re
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import load_workbook

from modes.trial_balance.tests.conftest_phase2 import SCREEN_ROWS, write_canonical

from modes.trial_balance.router import (  # noqa: E402
    _run_comparison_analytics_chain,
    _run_core_analytics_chain,
)
from modes.trial_balance.pipeline.tools import find_identifiers  # noqa: E402
from modes.trial_balance.pipeline.tools import read_manifest  # noqa: E402
from modes.trial_balance.pipeline.tools import build_comparison_report  # noqa: E402

_LAYER1 = json.dumps([
    {"rule": "TB-000", "rule_name": "Sign", "status": "PASS", "message": "positive = debit."},
    {"rule": "TB-026", "rule_name": "Scale", "status": "WARNING", "message": "Scale not declared."},
])


@pytest.fixture(scope="module")
def comparative_run(tmp_path_factory):
    """A full CY chain, a PY canonical TB, the comparison chain and both writers.

    PY differs from CY in three deliberate ways -- rebalanced accounts, one ledger
    absent, one ledger extra -- so structural delta and variance carry real content
    rather than degrading to empty sections that would pass vacuously.
    """
    run = tmp_path_factory.mktemp("comparative")
    cy, py, comp = run / "cy", run / "py", run / "comparison"
    for d in (cy, py, comp):
        d.mkdir(parents=True)

    py_rows = []
    for r in SCREEN_ROWS:
        if r["gl_code"] == "7001":       # suspense is new in CY
            continue
        q = dict(r)
        q["closing_balance"] = r["closing_balance"] * 0.8
        py_rows.append(q)
    py_rows.append({"gl_code": "6500", "gl_name": "Discontinued Product Line",
                    "closing_balance": 1_500_000.0, "main_head": "Expenses",
                    "sub_head_1": "Other expenses", "account_type": "Expense",
                    "mapped_status": "MAPPED"})

    cy_tb = write_canonical(cy / "canonical_tb.parquet", SCREEN_ROWS)
    write_canonical(py / "canonical_tb.parquet", py_rows)
    for d in (cy, py):
        (d / "layer1_results.json").write_text(_LAYER1, encoding="utf-8")

    with patch("modes.trial_balance.router.get_agent") as m:
        m.return_value.llm_client = None
        _run_core_analytics_chain(str(cy_tb), str(cy))
        _run_comparison_analytics_chain(comp, cy, py)
        result = build_comparison_report(output_dir=str(comp))

    return comp, result


def _doc_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _sheets(comp):
    wb = load_workbook(comp / "TB_Comparison_Audit.xlsx", read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


class TestComparativeReportRuns:
    def test_succeeds_and_writes_both_deliverables(self, comparative_run):
        comp, result = comparative_run
        assert result["execution_status"] == "SUCCESS", result.get("errors")
        assert (comp / "TB_Comparison_Report.docx").exists()
        assert (comp / "TB_Comparison_Audit.xlsx").exists()

    def test_document_saved_after_the_workbook(self, comparative_run):
        """The deferred save. Saved before the workbook, the appendix would be empty
        while the workbook was full -- and nothing would raise."""
        comp, _ = comparative_run
        manifest = read_manifest(comp)
        assert manifest["sheet_count"] > 0
        assert (f"contains {manifest['sheet_count']} schedule(s)"
                in _doc_text(comp / "TB_Comparison_Report.docx"))


class TestSectionStructure:
    @pytest.mark.parametrize("heading", [
        "1. Executive Audit Dashboard",
        "2. Scope, Inputs and Limitations",
        "3. TB Integrity and Reconciliation",
        "4. Financial-Statement Reconstruction",
        "5. Mapping and Classification Quality",
        "6. Materiality Assessment (CY basis)",
        "7. Comparative Analytical Review",
        "8. FSLI and Ledger-Level Movements (Current Year)",
        "9. Ratio and Relationship Analytics",
        "10. Account-Risk Scoring (Current Year)",
        "11. Assertion-Level Risk Assessment",
        "12. Sensitive and Public-Sector Accounts",
        "13. Estimates and Judgemental Balances",
        "14. Journal and Fraud-Risk Indicators",
        "15. Disclosure-Risk Indicators",
        "16. Going-Concern Indicators",
        "17. Consolidated Audit Findings",
        "18. Recommended Audit Procedures",
        "19. Evidence Request List",
        "20. Management and Statutory-Auditor Questions",
        "21. Prior-Issue Follow-Up",
        "22. Suggested Audit Plan and Prioritisation",
        "23. Data Lineage and Methodology",
        "24. Limitations & No-Opinion Statement",
        "25. Detailed Appendices and Full TB Schedules",
    ])
    def test_section_present(self, comparative_run, heading):
        comp, _ = comparative_run
        assert heading in _doc_text(comp / "TB_Comparison_Report.docx")

    def test_sections_are_in_ascending_order(self, comparative_run):
        """A renumbering that leaves one section out of position produces a document
        numbered 1, 2, 3, 6, 8 ... which reads as broken to an audit party."""
        comp, _ = comparative_run
        doc = Document(str(comp / "TB_Comparison_Report.docx"))
        nums = [int(re.match(r"(\d+)\.", p.text).group(1))
                for p in doc.paragraphs
                if p.style.name == "Heading 1" and re.match(r"\d+\.", p.text)]
        assert nums == sorted(nums), f"sections out of order: {nums}"
        assert nums == list(range(1, 26)), f"expected 1..25, got {nums}"

    def test_comparison_specific_content_is_present(self, comparative_run):
        """Section 7 is the comparative half. If the PY leg degraded these would be
        empty and every other assertion here would still pass."""
        comp, _ = comparative_run
        text = _doc_text(comp / "TB_Comparison_Report.docx")
        assert "Variance analysis" in text
        assert "PY closing to CY opening continuity" in text


class TestCrossReferenceInvariant:
    def test_every_sheet_is_named_in_the_document(self, comparative_run):
        """The requirement, on the comparative side. The first run of this caught two
        canonical-TB sheets created directly through wb.create_sheet(), which existed
        in the workbook but entered neither the manifest nor the document."""
        comp, _ = comparative_run
        text = _doc_text(comp / "TB_Comparison_Report.docx")
        missing = [n for n in _sheets(comp) if n not in text]
        assert not missing, (
            f"{len(missing)} sheet(s) exist in the comparative workbook but are never "
            f"named in the document: {missing}"
        )

    def test_manifest_matches_the_workbook(self, comparative_run):
        comp, _ = comparative_run
        manifest = read_manifest(comp)
        assert {s["name"] for s in manifest["sheets"]} == set(_sheets(comp))
        assert manifest["unregistered_sheets"] == []

    def test_both_periods_canonical_tbs_are_scheduled(self, comparative_run):
        """A comparative report must let a reader reach either period's full TB."""
        comp, _ = comparative_run
        names = _sheets(comp)
        assert "Canonical TB (CY)" in names and "Canonical TB (PY)" in names

    def test_comparison_only_schedules_are_present(self, comparative_run):
        comp, _ = comparative_run
        names = _sheets(comp)
        for sheet in ("Variance", "New Ledgers (CY)", "Removed Ledgers (PY)", "Sign Convention"):
            assert sheet in names, f"{sheet} missing from the comparative workbook"

    def test_cy_phase2_populations_are_carried_across(self, comparative_run):
        """The CY leg ran the full analytics chain, so its populations belong in this
        workbook rather than a second one the reader has to go and find."""
        comp, _ = comparative_run
        names = _sheets(comp)
        for sheet in ("Findings Register", "Evidence Requests", "Management Queries"):
            assert sheet in names, f"{sheet} not carried into the comparative workbook"


class TestSafety:
    def test_pii_is_masked(self, comparative_run):
        comp, _ = comparative_run
        found = find_identifiers(_doc_text(comp / "TB_Comparison_Report.docx"))
        assert not found, f"unmasked identifiers reached the comparative document: {found}"

    def test_unavailable_sections_state_why(self, comparative_run):
        comp, _ = comparative_run
        text = _doc_text(comp / "TB_Comparison_Report.docx")
        assert "No prior audit report" in text
        assert "Journal-level analysis is not performed" in text
