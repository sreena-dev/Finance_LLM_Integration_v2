"""Wave 8 remark #15: the Exception Register sheet must surface each row's composite
score and its per-rule weight breakdown -- both already computed by
build_exception_consolidator, just never rendered."""

from openpyxl import Workbook

from modes.trial_balance.pipeline.tools.comparison import _write_exception_register


def test_composite_score_and_breakdown_columns_render():
    wb = Workbook()
    ws = wb.active

    clusters = [
        {"cluster_id": "C1", "cluster_name": "Sign Reversal Cluster", "cluster_severity": "HIGH",
         "risk_themes": ["SIGN_REVERSAL"], "member_count": 1, "total_balance": 100.0},
    ]
    exceptions = [
        {
            "cluster_context": {"cluster_id": "C1"},
            "business_context": {"gl_code": "1001", "gl_name": "Suspense Account"},
            "financial_context": {"closing_balance": 100.0},
            "scoring": {"severity": "High", "composite_score": 45.0},
            "triggered_rules": [
                {"rule_id": "SIGN_REVERSAL", "description": "Balance flipped sign vs prior period.", "weight": 25.0},
                {"rule_id": "ABOVE_OM", "description": "Balance exceeds overall materiality.", "weight": 20.0},
            ],
        },
    ]

    next_row = _write_exception_register(ws, 1, clusters, exceptions)
    assert next_row > 1

    headers = [c.value for c in ws[1]]
    assert "Composite Score" in headers
    assert "Score Breakdown" in headers
    score_col = headers.index("Composite Score") + 1
    breakdown_col = headers.index("Score Breakdown") + 1

    account_row = ws[2]
    assert account_row[score_col - 1].value == 45.0
    assert "SIGN_REVERSAL +25" in account_row[breakdown_col - 1].value
    assert "ABOVE_OM +20" in account_row[breakdown_col - 1].value
